from __future__ import annotations

from pathlib import Path

import pytest
from docx import Document

from sloth_valet.extract import extract_file
from sloth_valet.models import ExtractionReason


def _build_pdf(pages: list[str]) -> bytes:
    page_count = len(pages)
    font_number = 3 + 2 * page_count
    objects: list[bytes] = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
    ]
    kids = " ".join(f"{3 + 2 * index} 0 R" for index in range(page_count))
    objects.append(f"<< /Type /Pages /Kids [{kids}] /Count {page_count} >>".encode())
    for index, page_text in enumerate(pages):
        page_number = 3 + 2 * index
        content_number = page_number + 1
        objects.append(
            (
                f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
                f"/Contents {content_number} 0 R "
                f"/Resources << /Font << /F1 {font_number} 0 R >> >> >>"
            ).encode()
        )
        escaped = page_text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
        stream = f"BT /F1 12 Tf 72 720 Td ({escaped}) Tj ET".encode()
        objects.append(
            f"<< /Length {len(stream)} >>\nstream\n".encode()
            + stream
            + b"\nendstream"
        )
    objects.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")

    output = bytearray(b"%PDF-1.4\n%\xe2\xe3\xcf\xd3\n")
    offsets: list[int] = []
    for number, body in enumerate(objects, start=1):
        offsets.append(len(output))
        output.extend(f"{number} 0 obj\n".encode())
        output.extend(body)
        output.extend(b"\nendobj\n")

    xref_offset = len(output)
    output.extend(f"xref\n0 {len(objects) + 1}\n".encode())
    output.extend(b"0000000000 65535 f \n")
    for offset in offsets:
        output.extend(f"{offset:010d} 00000 n \n".encode())
    output.extend(
        (
            f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\n"
            f"startxref\n{xref_offset}\n%%EOF\n"
        ).encode()
    )
    return bytes(output)


@pytest.mark.parametrize("suffix", [".txt", ".md", ".py", ".csv"])
def test_text_like_files_return_first_characters_without_controls(
    tmp_path: Path,
    suffix: str,
) -> None:
    source = tmp_path / f"sample{suffix}"
    content = "alpha\x00 beta\r\ngamma\u200b\n" + ("z" * 100)
    source.write_text(content, encoding="utf-8")

    result = extract_file(source, max_chars=25)

    assert result.reason is None
    assert result.text == "alpha beta\ngamma\nzzzzzzzzzzzz"[0:25]
    assert len(result.text) == 25
    assert "\x00" not in result.text
    assert "\r" not in result.text
    assert "\u200b" not in result.text
    assert result.metadata.name == source.name
    assert result.metadata.extension == suffix
    assert result.metadata.size == source.stat().st_size


def test_empty_file_returns_metadata_and_reason(tmp_path: Path) -> None:
    source = tmp_path / "empty.txt"
    source.touch()

    result = extract_file(source)

    assert result.text == ""
    assert result.reason == ExtractionReason.EMPTY_FILE
    assert result.metadata.size == 0
    assert result.metadata.mime_type.startswith("text/")


def test_pdf_uses_only_the_first_requested_pages(tmp_path: Path) -> None:
    source = tmp_path / "notes.pdf"
    source.write_bytes(_build_pdf(["Page one", "Page two", "Page three"]))

    result = extract_file(source, max_pdf_pages=2)

    assert result.reason is None
    assert "Page one" in result.text
    assert "Page two" in result.text
    assert "Page three" not in result.text
    assert result.metadata.mime_type == "application/pdf"


def test_corrupt_pdf_never_raises(tmp_path: Path) -> None:
    source = tmp_path / "broken.pdf"
    source.write_bytes(b"%PDF-1.4\nthis is not a valid pdf")

    result = extract_file(source)

    assert result.text == ""
    assert result.reason == ExtractionReason.CORRUPT
    assert "PDF" in (result.detail or "")
    assert result.metadata.extension == ".pdf"


def test_docx_returns_paragraphs_and_truncates(tmp_path: Path) -> None:
    source = tmp_path / "letter.docx"
    document = Document()
    document.add_paragraph("First paragraph")
    document.add_paragraph("")
    document.add_paragraph("Second paragraph")
    document.save(str(source))

    result = extract_file(source, max_chars=20)

    assert result.reason is None
    assert result.text == "First paragraph\n\nSec"
    assert result.metadata.extension == ".docx"
    assert "wordprocessingml" in result.metadata.mime_type


def test_corrupt_docx_never_raises(tmp_path: Path) -> None:
    source = tmp_path / "broken.docx"
    source.write_bytes(b"not a zip package")

    result = extract_file(source)

    assert result.text == ""
    assert result.reason == ExtractionReason.CORRUPT


def test_image_returns_metadata_only(tmp_path: Path) -> None:
    source = tmp_path / "screenshot.png"
    source.write_bytes(b"\x89PNG\r\n\x1a\n" + (b"\x00" * 32))

    result = extract_file(source)

    assert result.text == ""
    assert result.reason == ExtractionReason.IMAGE
    assert result.metadata.name == "screenshot.png"
    assert result.metadata.extension == ".png"
    assert result.metadata.size == source.stat().st_size
    assert result.metadata.mime_type == "image/png"


def test_unknown_binary_returns_metadata_only(tmp_path: Path) -> None:
    source = tmp_path / "blob.dat"
    source.write_bytes(b"\x00\x01\x02\x03" * 100)

    result = extract_file(source)

    assert result.text == ""
    assert result.reason == ExtractionReason.BINARY
    assert result.metadata.mime_type == "application/octet-stream"


def test_unknown_text_like_file_is_extracted(tmp_path: Path) -> None:
    source = tmp_path / "notes.weird"
    source.write_text("A readable note", encoding="utf-8")

    result = extract_file(source)

    assert result.reason is None
    assert result.text == "A readable note"
