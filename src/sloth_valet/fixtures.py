"""Deterministic benchmark fixtures generated from committed labels."""

from __future__ import annotations

import base64
import gzip
import io
import json
import os
import random
import tarfile
import zipfile
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Any

from docx import Document

from sloth_valet.config import CATEGORY_PATTERN

_LOREM = (
    "steady",
    "local",
    "quiet",
    "morning",
    "archive",
    "settled",
    "desk",
    "folder",
    "review",
    "careful",
    "offline",
    "simple",
)
_PNG_1X1 = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk"
    "YAAAAAYAAjCB0C8AAAAASUVORK5CYII="
)
_JPEG_1X1 = base64.b64decode(
    "/9j/4AAQSkZJRgABAQEASABIAAD/2wBDAP//////////////////////////////"
    "////////////////////////////////////////////////////2wBDAf//////"
    "////////////////////////////////////////////////////////////wAARCAAB"
    "AAEBAREA/8QAFQABAQAAAAAAAAAAAAAAAAAAAAX/xAAUEAEAAAAAAAAAAAAAAAAA"
    "AAAA/9oADAMBAAIQAxAAAAF//8QAFBABAAAAAAAAAAAAAAAAAAAAAP/aAAgBAQAB"
    "PwB//8QAFBEBAAAAAAAAAAAAAAAAAAAAAP/aAAgBAwEBPwB//8QAFBEBAAAAAAAA"
    "AAAAAAAAAAAAAP/aAAgBAgEBPwB//8QAFBABAAAAAAAAAAAAAAAAAAAAAP/aAAgB"
    "AQABPwB//9k="
)


@dataclass(frozen=True, slots=True)
class FixtureLabel:
    file: str
    category: str
    description: str


@dataclass(frozen=True, slots=True)
class FixtureManifest:
    seed: int
    fixtures: tuple[FixtureLabel, ...]


def load_fixture_manifest(path: Path) -> FixtureManifest:
    """Load and strictly validate the committed benchmark labels."""
    try:
        raw = json.loads(path.read_text(encoding="utf-8"))
    except OSError as exc:
        raise OSError(f"cannot read labels file {path}: {exc}") from exc
    except json.JSONDecodeError as exc:
        raise ValueError(f"labels file {path} is not valid JSON: {exc}") from exc
    if not isinstance(raw, dict):
        raise ValueError("labels file must be a JSON object")
    seed = raw.get("seed")
    if not isinstance(seed, int) or isinstance(seed, bool):
        raise ValueError("labels seed must be an integer")
    entries = raw.get("fixtures")
    if not isinstance(entries, list) or not entries:
        raise ValueError("labels fixtures must be a non-empty list")

    labels: list[FixtureLabel] = []
    names: set[str] = set()
    for index, entry in enumerate(entries):
        if not isinstance(entry, dict):
            raise ValueError(f"fixture {index} must be an object")
        if set(entry) != {"file", "category", "description"}:
            raise ValueError(
                f"fixture {index} must contain exactly file, category, and description"
            )
        file = _safe_filename(entry["file"])
        category = entry["category"]
        description = entry["description"]
        if not isinstance(category, str) or not CATEGORY_PATTERN.fullmatch(category):
            raise ValueError(f"fixture {file!r} has an invalid category")
        if not isinstance(description, str) or not description.strip():
            raise ValueError(f"fixture {file!r} has an empty description")
        if file in names:
            raise ValueError(f"duplicate fixture name {file!r}")
        names.add(file)
        labels.append(FixtureLabel(file=file, category=category, description=description))
    return FixtureManifest(seed=seed, fixtures=tuple(labels))


def generate_fixtures(
    manifest: FixtureManifest,
    output_dir: Path,
    *,
    overwrite: bool = False,
) -> tuple[Path, ...]:
    """Create every labeled fixture and return the written paths."""
    output_dir.mkdir(parents=True, exist_ok=True)
    if not output_dir.is_dir():
        raise OSError(f"fixture output is not a directory: {output_dir}")

    expected = [output_dir / label.file for label in manifest.fixtures]
    existing = [path for path in expected if os.path.lexists(path)]
    if existing and not overwrite:
        raise FileExistsError(
            f"{len(existing)} fixture file(s) already exist; pass --overwrite to regenerate"
        )

    written: list[Path] = []
    for label in manifest.fixtures:
        path = output_dir / label.file
        if path.is_symlink():
            raise OSError(f"refusing to replace symlinked fixture: {path}")
        data = _fixture_bytes(label, manifest.seed)
        _atomic_write(path, data)
        written.append(path)
    return tuple(written)


def _fixture_bytes(label: FixtureLabel, seed: int) -> bytes:
    name = label.file
    lower = name.casefold()
    if lower.endswith(".pdf"):
        return _minimal_pdf(label.description + "\n\n" + _filler(seed, label))
    if lower.endswith(".docx"):
        return _docx_bytes(label)
    if lower.endswith(".tar.gz"):
        return _tar_gz_bytes(_archive_readme(label))
    if lower.endswith(".zip"):
        return _zip_bytes(_archive_readme(label))
    if lower.endswith(".png"):
        return _PNG_1X1
    if lower.endswith(".jpg") or lower.endswith(".jpeg"):
        return _JPEG_1X1
    if lower.endswith(".mp3"):
        return b"ID3\x04\x00\x00\x00\x00\x00\x00" + _text_bytes(label, seed)
    if lower.endswith(".mp4") or lower.endswith(".m4a"):
        return b"\x00\x00\x00\x18ftypM4A \x00\x00\x00\x00M4A mp42" + _text_bytes(
            label, seed
        )
    if lower.endswith(".deb"):
        return b"!<arch>\ndebian-binary   " + _text_bytes(label, seed)
    if lower.endswith(".appimage"):
        return b"\x7fELF" + _text_bytes(label, seed)
    if lower.endswith(".iso"):
        return b"CD001" + _text_bytes(label, seed)
    if lower.endswith(".exe"):
        return b"MZ" + _text_bytes(label, seed)
    if lower.endswith(".dat"):
        return b"\x00\x01\x02\x03" + _text_bytes(label, seed)
    return _text_bytes(label, seed)


def _text_bytes(label: FixtureLabel, seed: int) -> bytes:
    suffix = Path(label.file).suffix.casefold()
    description = label.description
    filler = _filler(seed, label)
    if suffix == ".csv":
        rows = [
            "item,detail,amount",
            f"1,{description.replace(',', ';')},",
            f"2,{filler.replace(',', ';')},",
        ]
        return ("\n".join(rows) + "\n").encode("utf-8")
    if suffix == ".json":
        return (
            json.dumps(
                {
                    "category": label.category,
                    "description": description,
                    "notes": filler,
                    "items": [
                        {"name": word, "value": index}
                        for index, word in enumerate(_LOREM[:4])
                    ],
                },
                indent=2,
                sort_keys=True,
            )
            + "\n"
        ).encode("utf-8")
    if suffix == ".py":
        return (
            f'"""{description}\n\n{filler}\n"""\n\n'
            "from __future__ import annotations\n\n\n"
            "def main() -> None:\n"
            f'    print("{description}")\n\n\n'
            'if __name__ == "__main__":\n'
            "    main()\n"
        ).encode()
    if suffix == ".sh":
        return (
            "#!/bin/sh\n"
            f"# {description}\n"
            "set -eu\n"
            f'echo "{filler}"\n'
        ).encode()
    if suffix == ".md":
        return (
            f"# {Path(label.file).stem.replace('-', ' ').title()}\n\n"
            f"{description}\n\n{filler}\n"
        ).encode()
    return f"{description}\n\n{filler}\n".encode()


def _docx_bytes(label: FixtureLabel) -> bytes:
    document = Document()
    document.core_properties.title = Path(label.file).stem
    document.core_properties.subject = label.category
    document.core_properties.created = datetime(2026, 10, 20, 3, 0, 0)
    document.core_properties.modified = datetime(2026, 10, 20, 3, 0, 0)
    document.add_heading(Path(label.file).stem.replace("-", " ").title(), level=1)
    document.add_paragraph(label.description)
    document.add_paragraph(
        "This fixture was generated deterministically for the Sloth Valet benchmark."
    )
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _archive_readme(label: FixtureLabel) -> bytes:
    return (
        f"{label.description}\n\nCategory: {label.category}\n"
        "Generated fixture for the Sloth Valet benchmark.\n"
    ).encode()


def _zip_bytes(readme: bytes) -> bytes:
    buffer = io.BytesIO()
    info = zipfile.ZipInfo("README.txt", date_time=(2026, 10, 20, 3, 0, 0))
    info.compress_type = zipfile.ZIP_DEFLATED
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr(info, readme)
    return buffer.getvalue()


def _tar_gz_bytes(readme: bytes) -> bytes:
    buffer = io.BytesIO()
    with (
        gzip.GzipFile(fileobj=buffer, mode="wb", mtime=0) as compressed,
        tarfile.open(fileobj=compressed, mode="w") as archive,
    ):
        info = tarfile.TarInfo("README.txt")
        info.size = len(readme)
        info.mtime = 0
        archive.addfile(info, io.BytesIO(readme))
    return buffer.getvalue()


def _minimal_pdf(text: str) -> bytes:
    lines = _wrap_text(text, 82)
    escaped = [
        line.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
        for line in lines
    ]
    stream = (
        "BT\n/F1 12 Tf\n14 TL\n50 750 Td\n"
        + "".join(f"({line}) Tj T*\n" for line in escaped)
        + "ET\n"
    ).encode("latin-1", "replace")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        (
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            b"/Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>"
        ),
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"endstream",
    ]
    output = bytearray(b"%PDF-1.4\n%\xe2\xe3\xcf\xd3\n")
    offsets = [0]
    for number, obj in enumerate(objects, start=1):
        offsets.append(len(output))
        output.extend(f"{number} 0 obj\n".encode())
        output.extend(obj)
        output.extend(b"\nendobj\n")
    xref_offset = len(output)
    output.extend(f"xref\n0 {len(objects) + 1}\n".encode())
    output.extend(b"0000000000 65535 f \n")
    for offset in offsets[1:]:
        output.extend(f"{offset:010d} 00000 n \n".encode())
    output.extend(
        (
            f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\n"
            f"startxref\n{xref_offset}\n%%EOF\n"
        ).encode()
    )
    return bytes(output)


def _wrap_text(text: str, width: int) -> list[str]:
    lines: list[str] = []
    for paragraph in text.splitlines() or [""]:
        words = paragraph.split()
        current = ""
        for word in words:
            candidate = f"{current} {word}".strip()
            if len(candidate) > width and current:
                lines.append(current)
                current = word
            else:
                current = candidate
        if current:
            lines.append(current)
    return lines or [""]


def _filler(seed: int, label: FixtureLabel) -> str:
    generator = random.Random(f"{seed}:{label.file}")
    words = [generator.choice(_LOREM) for _ in range(18)]
    first = " ".join(words[:9]).capitalize()
    second = " ".join(words[9:]).capitalize()
    return f"{first}. {second}."


def _safe_filename(value: Any) -> str:
    if not isinstance(value, str) or not value:
        raise ValueError("fixture file must be a non-empty string")
    if (
        value != value.strip()
        or value in {".", ".."}
        or value.startswith(".")
        or "/" in value
        or "\\" in value
        or "\x00" in value
    ):
        raise ValueError(f"unsafe fixture filename {value!r}")
    return value


def _atomic_write(path: Path, data: bytes) -> None:
    temporary = path.with_name(f".{path.name}.tmp")
    try:
        with temporary.open("xb") as output:
            output.write(data)
            output.flush()
            os.fsync(output.fileno())
        os.replace(temporary, path)
    finally:
        if temporary.exists():
            temporary.unlink()
